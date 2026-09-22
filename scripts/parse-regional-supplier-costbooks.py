#!/usr/bin/env python3
"""Parse normalized TradeOS supplier workbooks into reviewable intake JSON.

This script stages evidence only. It does not connect to Postgres, change
Material.unitCost, or create SupplierPriceUpdate review proposals.
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - operator environment check
    raise SystemExit("Install openpyxl to parse XLSX inputs: python -m pip install openpyxl") from exc


STATUS_MAP = {
    "VERIFIED_CURRENT": "priced",
    "VERIFIED_REGIONAL": "priced",
    "UNAVAILABLE": "unavailable",
    "NOT_LISTED": "not-listed",
    "NOT_AVAILABLE": "unavailable",
}


def cell(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip()


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Expected numeric value, got {value!r}") from exc


def boolean(value: Any, default: bool = False) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "yes", "1", "active", "available"}


def confidence(value: Any) -> str | None:
    score = number(value)
    if score is None:
        return None
    return "high" if score >= 90 else "medium" if score >= 70 else "low"


def observed_at(value: Any) -> str:
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc).isoformat().replace("+00:00", "Z")
    raw = text(value)
    if not raw:
        raise ValueError("Current_Prices.observed_at is required")
    parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    dt = parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing required sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    except StopIteration as exc:
        raise ValueError(f"Sheet {sheet_name!r} is empty") from exc
    return [dict(zip(headers, values)) for values in iterator if any(value not in (None, "") for value in values)]


def parse_workbook(source_name: str, workbook_bytes: bytes) -> dict[str, Any]:
    from io import BytesIO

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    manifest = rows(workbook, "Import_Manifest")
    manifest_map = {text(row.get("parameter") or row.get("field")): text(row.get("value")) for row in manifest}
    products = rows(workbook, "Supplier_Products")
    prices = rows(workbook, "Current_Prices")

    supplier_code = text(manifest_map.get("supplier_code")) or text(cell(products[0], "supplier_code"))
    supplier_name = text(manifest_map.get("supplier_name")) or text(cell(products[0], "supplier_name"))
    if not supplier_code or not supplier_name:
        raise ValueError(f"{source_name}: supplier_code and supplier_name are required")

    product_rows = []
    product_keys = set()
    for row in products:
        key = text(cell(row, "supplier_product_key"))
        name = text(cell(row, "supplier_product_name", "product_name"))
        if not key or not name:
            raise ValueError(f"{source_name}: every Supplier_Products row needs a product key and name")
        if key in product_keys:
            raise ValueError(f"{source_name}: duplicate supplier product key {key}")
        product_keys.add(key)
        availability = text(cell(row, "availability"))
        if availability is None:
            availability_status = "unknown"
        else:
            normalized_availability = availability.upper().replace("-", "_").replace(" ", "_")
            if normalized_availability == "AVAILABLE":
                availability_status = "available"
            elif normalized_availability in {"UNAVAILABLE", "NOT_AVAILABLE", "NOT_LISTED"}:
                availability_status = "unavailable"
            else:
                raise ValueError(f"{source_name}: unsupported availability {availability!r}")
        product_rows.append({
            "supplierProductKey": key,
            "canonicalMaterialKey": text(cell(row, "canonical_material_key")),
            "sku": text(cell(row, "supplier_sku", "sku")),
            "manufacturerPartNumber": text(cell(row, "model_number", "manufacturer_part_number")),
            "name": name,
            "packageDescription": text(cell(row, "package_description")),
            "packageQuantity": number(cell(row, "package_quantity")),
            "purchaseUnit": text(cell(row, "purchase_unit")),
            "productUrl": text(cell(row, "product_url")),
            "availabilityStatus": availability_status,
            "isActive": boolean(cell(row, "active"), True),
            "sourceFile": source_name,
        })

    observation_rows = []
    observation_keys = set()
    for excel_row, row in enumerate(prices, start=2):
        key = text(cell(row, "price_observation_id")) or f"{source_name}:{excel_row}"
        product_key = text(cell(row, "supplier_product_key"))
        if not product_key or product_key not in product_keys:
            raise ValueError(f"{source_name}: price row {excel_row} references unknown supplier product {product_key!r}")
        if key in observation_keys:
            raise ValueError(f"{source_name}: duplicate price observation key {key}")
        observation_keys.add(key)
        raw_status = (text(cell(row, "price_status")) or "UNAVAILABLE").upper().replace(" ", "_")
        status = STATUS_MAP.get(raw_status)
        if status is None:
            raise ValueError(f"{source_name}: unsupported price_status {raw_status!r}")
        observation_rows.append({
            "observationKey": key,
            "supplierProductKey": product_key,
            "marketCode": text(cell(row, "market_code")),
            "storeName": text(cell(row, "store_name")),
            "city": text(cell(row, "city")),
            "state": text(cell(row, "state")),
            "postalCode": text(cell(row, "zip_code", "postal_code")),
            "observedAt": observed_at(cell(row, "observed_at")),
            "sourceFile": source_name,
            "sourceRow": excel_row,
            "currency": text(cell(row, "currency")) or "USD",
            "priceStatus": status,
            "regularPrice": number(cell(row, "regular_price")),
            "salePrice": number(cell(row, "sale_price")),
            "rebatePrice": number(cell(row, "rebate_price")),
            "effectivePrice": number(cell(row, "effective_purchase_price", "effective_price")),
            "purchaseUnit": text(cell(row, "purchase_unit")),
            "packageQuantity": number(cell(row, "package_quantity")),
            "normalizedUnitPrice": number(cell(row, "normalized_unit_price")),
            "normalizedUnit": text(cell(row, "normalized_unit")),
            "sourceUrl": text(cell(row, "source_url")),
            "eligibilityReason": text(cell(row, "eligibility_reason")),
            "sourceConfidence": confidence(cell(row, "confidence_score", "source_confidence", "confidence")),
        })

    return {
        "schemaVersion": manifest_map.get("schema_version") or "TRADEOS_COSTBOOK_V1",
        "sourceFile": source_name,
        "supplierCode": supplier_code,
        "supplierName": supplier_name,
        "products": product_rows,
        "observations": observation_rows,
    }


def parse_path(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".zip":
        with ZipFile(path) as archive:
            return [parse_workbook(name, archive.read(name)) for name in sorted(archive.namelist()) if name.lower().endswith(".xlsx")]
    return [parse_workbook(path.name, path.read_bytes())]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="An XLSX workbook or ZIP containing normalized supplier workbooks")
    parser.add_argument("-o", "--output", type=Path, required=True, help="Output JSON intake-batch path")
    args = parser.parse_args()
    batches = parse_path(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps({"batches": batches}, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "workbooks": len(batches),
        "products": sum(len(batch["products"]) for batch in batches),
        "observations": sum(len(batch["observations"]) for batch in batches),
        "unavailable": sum(sum(row["priceStatus"] != "priced" for row in batch["observations"]) for batch in batches),
        "output": str(args.output),
    }))


if __name__ == "__main__":
    main()
